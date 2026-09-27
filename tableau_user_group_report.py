"""
Tableau User & Group Report

Connects to a Tableau Cloud or Tableau Server site using a Personal Access
Token (PAT), collects every user and every group on the site, and writes an
Excel workbook with three sheets:

    Stats          Site details, headline counts and a site role breakdown
    User Summary   One row per user, with all the groups they belong to
    Group Details  One row per group member (empty groups are listed too)

Connection details are read from config.json (see config.example.json).
"""

import argparse
import json
import os
import sys
import urllib.error
import urllib.request
from collections import Counter
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

try:
    from zoneinfo import ZoneInfo
except ImportError:  # Python older than 3.9
    ZoneInfo = None


# --- Settings ----------------------------------------------------------------
DEFAULT_CONFIG_FILE = "config.json"
DEFAULT_API_VERSION = "3.21"
DEFAULT_OUTPUT_DIR  = "output"
PAGE_SIZE           = 100

# --- Colour palette ----------------------------------------------------------
HDR_FILL_DARK  = PatternFill("solid", fgColor="1F3864")
HDR_FILL_MED   = PatternFill("solid", fgColor="2E75B6")
HDR_FILL_LIGHT = PatternFill("solid", fgColor="D6E4F0")
ROW_ALT_FILL   = PatternFill("solid", fgColor="EBF3FB")
NO_GROUP_FILL  = PatternFill("solid", fgColor="FFF2CC")
WHITE_FILL     = PatternFill("solid", fgColor="FFFFFF")

THIN_BORDER = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)


# --- Configuration -----------------------------------------------------------
def load_config(path):
    """Read and validate the confidential config file."""
    if not os.path.exists(path):
        sys.exit(
            f"ERROR: Config file '{path}' not found.\n"
            "Copy config.example.json to config.json and fill in your details."
        )

    with open(path, "r", encoding="utf-8") as f:
        cfg = json.load(f)

    required = ["server_url", "site_content_url", "pat_name", "pat_secret"]
    missing = [k for k in required if not str(cfg.get(k, "")).strip()]
    if missing:
        sys.exit(f"ERROR: Missing values in {path}: {', '.join(missing)}")

    cfg["server_url"]  = cfg["server_url"].rstrip("/")
    cfg["api_version"] = cfg.get("api_version") or DEFAULT_API_VERSION
    cfg["output_dir"]  = cfg.get("output_dir") or DEFAULT_OUTPUT_DIR
    cfg["timezone"]    = cfg.get("timezone") or ""
    return cfg


def now_in_timezone(tz_name):
    """Current time in the configured timezone, or local machine time."""
    if tz_name and ZoneInfo:
        try:
            return datetime.now(ZoneInfo(tz_name))
        except Exception:
            print(f"  ! Unknown timezone '{tz_name}', using local machine time.")
    return datetime.now()


# --- REST helper -------------------------------------------------------------
def rest(url, method="GET", data=None, headers=None):
    h = {"Content-Type": "application/json", "Accept": "application/json"}
    if headers:
        h.update(headers)
    body = json.dumps(data).encode() if data else None
    req = urllib.request.Request(url, data=body, headers=h, method=method)
    try:
        with urllib.request.urlopen(req, timeout=30) as resp:
            raw = resp.read()
            return json.loads(raw) if raw else {}
    except urllib.error.HTTPError as e:
        raise RuntimeError(f"HTTP {e.code} on {url}\n{e.read().decode()}")


# --- Tableau calls -----------------------------------------------------------
def sign_in(cfg):
    print("  Signing in ...")
    payload = {
        "credentials": {
            "personalAccessTokenName":   cfg["pat_name"],
            "personalAccessTokenSecret": cfg["pat_secret"],
            "site": {"contentUrl": cfg["site_content_url"]},
        }
    }
    url  = f"{cfg['server_url']}/api/{cfg['api_version']}/auth/signin"
    resp = rest(url, "POST", payload)
    token   = resp["credentials"]["token"]
    site_id = resp["credentials"]["site"]["id"]
    print("  Signed in")
    return token, site_id


def sign_out(cfg, token):
    try:
        rest(f"{cfg['server_url']}/api/{cfg['api_version']}/auth/signout",
             "POST", headers={"X-Tableau-Auth": token})
        print("  Signed out")
    except Exception:
        pass  # signing out is a courtesy; the token expires anyway


def paginate(cfg, site_id, token, endpoint, collection_key, item_key):
    """Fetch every page of a Tableau REST list endpoint."""
    results, page = [], 1
    while True:
        url = (f"{cfg['server_url']}/api/{cfg['api_version']}/sites/{site_id}/"
               f"{endpoint}?pageSize={PAGE_SIZE}&pageNumber={page}")
        resp  = rest(url, headers={"X-Tableau-Auth": token})
        items = resp.get(collection_key, {}).get(item_key, [])
        if isinstance(items, dict):
            items = [items]
        results.extend(items)
        total = int(resp.get("pagination", {}).get("totalAvailable", len(results)))
        if page * PAGE_SIZE >= total:
            break
        page += 1
    return results


def new_user_record(u):
    return {
        "id":        u["id"],
        "name":      u.get("name", ""),
        "fullName":  u.get("fullName", u.get("name", "")),
        "email":     u.get("email", ""),
        "role":      u.get("siteRole", ""),
        "lastLogin": u.get("lastLogin", ""),
        "groups":    [],
    }


def collect_data(cfg, site_id, token):
    """Return users (with their groups) and groups (with their members)."""
    print("  Fetching users ...")
    users = paginate(cfg, site_id, token, "users", "users", "user")
    print(f"  {len(users)} users found")

    print("  Fetching groups ...")
    groups = paginate(cfg, site_id, token, "groups", "groups", "group")
    print(f"  {len(groups)} groups found")

    user_map = {u["id"]: new_user_record(u) for u in users}
    group_members = {}

    for g in groups:
        gname = g.get("name", "")
        print(f"    {gname}: ", end="", flush=True)
        members = paginate(cfg, site_id, token,
                           f"groups/{g['id']}/users", "users", "user")
        print(f"{len(members)} members")

        group_members[gname] = []
        for m in members:
            if m["id"] not in user_map:
                user_map[m["id"]] = new_user_record(m)
            user_map[m["id"]]["groups"].append(gname)
            group_members[gname].append(user_map[m["id"]])

    return user_map, group_members


# --- Excel helpers -----------------------------------------------------------
def write_cell(ws, row, col, value, font=None, fill=None, alignment=None, border=None):
    c = ws.cell(row=row, column=col, value=value)
    if font:      c.font      = font
    if fill:      c.fill      = fill
    if alignment: c.alignment = alignment
    if border:    c.border    = border
    return c


def write_header_row(ws, row, headers, fill):
    font  = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col, text in enumerate(headers, 1):
        write_cell(ws, row, col, text, font, fill, align, THIN_BORDER)
    ws.row_dimensions[row].height = 22


def auto_width(ws, min_w=10, max_w=55):
    for col in ws.columns:
        longest = max((len(str(c.value)) for c in col if c.value is not None), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = \
            min(max(longest + 2, min_w), max_w)


def short_date(value):
    return value[:10] if value else ""


# --- Sheet 1: Stats ----------------------------------------------------------
def build_stats_sheet(wb, cfg, user_map, group_members, generated_at):
    ws = wb.create_sheet("Stats")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 40

    title = ws.cell(row=1, column=1, value="Tableau User & Group Report")
    title.font = Font(name="Arial", bold=True, size=16, color="1F3864")
    ws.merge_cells("A1:B1")
    ws.row_dimensions[1].height = 32

    users = list(user_map.values())
    rows = [
        ("Server",              cfg["server_url"]),
        ("Site",                cfg["site_content_url"]),
        ("Generated on",        generated_at.strftime("%Y-%m-%d %H:%M:%S")),
        ("", ""),
        ("Total Users",         len(users)),
        ("Total Groups",        len(group_members)),
        ("Empty Groups",        sum(1 for m in group_members.values() if not m)),
        ("Users with no group", sum(1 for u in users if not u["groups"])),
        ("Users in 2+ groups",  sum(1 for u in users if len(u["groups"]) >= 2)),
    ]

    key_font = Font(name="Arial", bold=True, size=10, color="1F3864")
    val_font = Font(name="Arial", size=10)
    r = 3
    for key, val in rows:
        if key:
            write_cell(ws, r, 1, key, key_font, alignment=Alignment(horizontal="left"))
            c = write_cell(ws, r, 2, val, val_font, alignment=Alignment(horizontal="right"))
            if isinstance(val, int):
                c.fill = HDR_FILL_LIGHT
        r += 1

    r += 1
    write_header_row(ws, r, ["Site Role", "Users"], HDR_FILL_DARK)
    r += 1
    small = Font(name="Arial", size=9)
    for role, count in sorted(Counter(u["role"] for u in users).items(),
                              key=lambda x: -x[1]):
        write_cell(ws, r, 1, role, small, border=THIN_BORDER)
        write_cell(ws, r, 2, count, small, alignment=Alignment(horizontal="right"),
                   border=THIN_BORDER)
        r += 1


# --- Sheet 2: User Summary ---------------------------------------------------
def build_user_summary_sheet(wb, user_map):
    ws = wb.create_sheet("User Summary")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"

    headers = ["#", "Display Name", "Username", "Email", "Site Role", "Groups", "Last Login"]
    write_header_row(ws, 1, headers, HDR_FILL_DARK)

    font  = Font(name="Arial", size=9)
    align = Alignment(vertical="center", wrap_text=True)

    for i, u in enumerate(sorted(user_map.values(), key=lambda x: x["name"].lower()), 1):
        groups_str = ", ".join(sorted(u["groups"])) if u["groups"] else "(none)"
        fill = NO_GROUP_FILL if not u["groups"] else (ROW_ALT_FILL if i % 2 == 0 else WHITE_FILL)
        values = [i, u["fullName"] or u["name"], u["name"], u["email"],
                  u["role"], groups_str, short_date(u["lastLogin"])]
        for col, val in enumerate(values, 1):
            write_cell(ws, i + 1, col, val, font, fill, align, THIN_BORDER)

    auto_width(ws)
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["F"].width = 45


# --- Sheet 3: Group Details --------------------------------------------------
def build_group_details_sheet(wb, group_members):
    ws = wb.create_sheet("Group Details")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"

    headers = ["#", "Group Name", "Display Name", "Username", "Email", "Site Role", "Last Login"]
    write_header_row(ws, 1, headers, HDR_FILL_MED)

    font      = Font(name="Arial", size=9)
    bold_font = Font(name="Arial", size=9, bold=True)
    align     = Alignment(vertical="center")

    row = 2
    for gname in sorted(group_members, key=str.lower):
        members = sorted(group_members[gname], key=lambda x: x["name"].lower())
        if not members:
            members = [{"fullName": "(empty group)", "name": "", "email": "",
                        "role": "", "lastLogin": ""}]
        for u in members:
            n = row - 1
            fill = ROW_ALT_FILL if n % 2 == 0 else WHITE_FILL
            values = [n, gname, u.get("fullName") or u.get("name", ""), u.get("name", ""),
                      u.get("email", ""), u.get("role", ""), short_date(u.get("lastLogin"))]
            for col, val in enumerate(values, 1):
                write_cell(ws, row, col, val, bold_font if col == 2 else font,
                           fill, align, THIN_BORDER)
            row += 1

    auto_width(ws)
    ws.column_dimensions["A"].width = 6


# --- Main --------------------------------------------------------------------
def main():
    parser = argparse.ArgumentParser(
        description="Export Tableau users and groups to an Excel report.")
    parser.add_argument("--config", default=DEFAULT_CONFIG_FILE,
                        help="Path to the config file (default: config.json)")
    args = parser.parse_args()

    cfg = load_config(args.config)
    generated_at = now_in_timezone(cfg["timezone"])

    os.makedirs(cfg["output_dir"], exist_ok=True)
    output_file = os.path.join(
        cfg["output_dir"],
        f"Tableau_User_Group_Report_{generated_at.strftime('%Y%m%d_%H%M%S')}.xlsx")

    print("\n" + "=" * 55)
    print("  Tableau User & Group Report")
    print(f"  Server : {cfg['server_url']}")
    print(f"  Site   : {cfg['site_content_url']}")
    print("=" * 55 + "\n")

    token, site_id = sign_in(cfg)
    try:
        print("\nFetching data ...")
        user_map, group_members = collect_data(cfg, site_id, token)
    finally:
        sign_out(cfg, token)

    print("\nBuilding Excel workbook ...")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    build_stats_sheet(wb, cfg, user_map, group_members, generated_at)
    build_user_summary_sheet(wb, user_map)
    build_group_details_sheet(wb, group_members)
    wb.save(output_file)

    print(f"\nDone. Report saved to: {output_file}")
    print(f"  Users  : {len(user_map)}")
    print(f"  Groups : {len(group_members)}\n")


if __name__ == "__main__":
    main()
